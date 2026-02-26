"""Test handler sync debugging"""
import sys
import io
from pathlib import Path

# Fix encoding for Windows console
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

sys.path.insert(0, 'src')

from core.database import get_database
from core.excel_handler import get_excel_handler
from utils.config import config

def main():
    db = get_database(config.db_path)
    excel_handler = get_excel_handler()
    
    print("=== Departments ===")
    depts = db.get_departments()
    for d in depts:
        print(f"ID: {d['id']}, Name: {d['name']}")
        print(f"  Path: {d['network_path']}")
        p = Path(d['network_path'])
        print(f"  Exists: {p.exists()}, IsDir: {p.is_dir() if p.exists() else 'N/A'}")
        
        if p.exists() and p.is_dir():
            for f in p.glob('File tổng hợp*.xls*'):
                print(f"  Found file: {f.name}")
                
                # Check for DS KDTVN sheet
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(str(f), read_only=True)
                    print(f"  Sheets: {wb.sheetnames}")
                    if 'DS KDTVN' in wb.sheetnames:
                        print("  ✅ Found DS KDTVN sheet!")
                        # Try reading handlers
                        handlers = excel_handler.read_handlers(str(f))
                        print(f"  Handlers found: {len(handlers)}")
                        if handlers:
                            for h in handlers[:5]:
                                print(f"    - {h['name']}: {h['email']}")
                    wb.close()
                except Exception as e:
                    print(f"  ❌ Error: {e}")
        print()
    
    print("\n=== Current Handlers in DB ===")
    handlers = db.get_handlers()
    print(f"Total: {len(handlers)}")
    for h in handlers:
        print(f"  {h['id']}: {repr(h['name'])} - {h.get('email')}")

if __name__ == "__main__":
    main()
