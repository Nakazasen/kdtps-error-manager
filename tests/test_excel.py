"""
Unit Tests for ExcelHandler
"""
import unittest
import tempfile
import os
from pathlib import Path

# Add src to path for imports
import sys
sys.path.insert(0, str(Path(__file__).parent.parent / 'src'))

from openpyxl import Workbook
from core.excel_handler import ExcelHandler, get_excel_handler


class TestExcelHandler(unittest.TestCase):
    """Test ExcelHandler functionality."""
    
    def setUp(self):
        """Create temporary test Excel files."""
        self.temp_dir = tempfile.mkdtemp()
        self.handler = ExcelHandler()
        
        # Create test KDTPS file
        self.kdtps_file = Path(self.temp_dir) / "test_kdtps.xlsx"
        self._create_test_kdtps_file()
        
        # Create test summary file
        self.summary_file = Path(self.temp_dir) / "test_summary.xlsx"
    
    def tearDown(self):
        """Clean up temporary files."""
        for f in Path(self.temp_dir).glob("*"):
            os.remove(f)
        os.rmdir(self.temp_dir)
    
    def _create_test_kdtps_file(self):
        """Create a test KDTPS Excel file."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Máy in"
        
        # Row 1 - empty or title
        ws.cell(row=1, column=1, value="KDTPS Report")
        
        # Row 2 - Headers (KDTPS format)
        headers = ["No.", "Ngày", "Loại máy", "Line", "Số lượng", 
                   "F", "Cxxx", "Jxxx", "Mã lỗi", "Nội dung", 
                   "K", "L", "M", "Điều tra", "O", "P"]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=2, column=col_idx, value=header)
        
        # Row 3+ - Data
        data_rows = [
            ["001", "2026-01-15", "MachineA", "Line01", 5, "", "C001", "J001", "E001", "Test error 1", "", "", "", "Investigation 1", "", ""],
            ["002", "2026-01-16", "MachineB", "Line02", 3, "", "C002", "J002", "E002", "Test error 2", "", "", "", "Investigation 2", "", ""],
            ["003", "2026-01-17", "MachineA", "Line01", 2, "", "C003", "", "E003", "Test error 3", "", "", "", "", "", ""],
        ]
        
        for row_idx, row_data in enumerate(data_rows, start=3):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        wb.save(self.kdtps_file)
        wb.close()
    
    def test_read_kdtps_file(self):
        """Test reading KDTPS file with correct row offsets."""
        headers, data = self.handler.read_kdtps_file(
            str(self.kdtps_file),
            sheet_name="Máy in",
            start_row=3,  # Data starts at row 3
            header_row=2  # Header at row 2
        )
        
        # Check headers
        self.assertEqual(len(headers), 16)  # A-P = 16 columns
        self.assertEqual(headers[0], "No.")
        
        # Check data
        self.assertEqual(len(data), 3)  # 3 data rows
        self.assertEqual(data[0][0], "001")
        self.assertEqual(data[1][2], "MachineB")
    
    def test_get_sheet_names(self):
        """Test getting sheet names."""
        sheets = self.handler.get_sheet_names(str(self.kdtps_file))
        self.assertIn("Máy in", sheets)
    
    def test_filter_rows_by_line(self):
        """Test filtering rows by production line."""
        _, data = self.handler.read_kdtps_file(str(self.kdtps_file), start_row=3, header_row=2)
        
        # Filter for Line01 only (column D = index 3)
        filtered = self.handler.filter_rows_by_line(data, 3, ["Line01"])
        
        self.assertEqual(len(filtered), 2)  # Should have 2 records for Line01
        self.assertTrue(all(row[3] == "Line01" for row in filtered))
    
    def test_write_to_summary_file_new(self):
        """Test writing to new summary file."""
        # Prepare data
        data_rows = [
            ["001", "2026-01-15", "MachineA", "Line01", 5, "", "C001", "J001", "E001", "Error 1", "", "", "", "Inv 1", "", ""],
        ]
        
        stats = self.handler.write_to_summary_file(
            str(self.summary_file),
            "Máy in",
            data_rows,
            start_row=2,
            key_column_idx=0
        )
        
        self.assertEqual(stats['inserted'], 1)
        self.assertEqual(stats['updated'], 0)
        self.assertTrue(self.summary_file.exists())
    
    def test_write_to_summary_file_update(self):
        """Test updating existing records in summary file."""
        # Write initial data
        data_rows_1 = [
            ["001", "2026-01-15", "MachineA", "Line01", 5, "", "C001", "", "", "", "", "", "", "Original", "", ""],
        ]
        self.handler.write_to_summary_file(str(self.summary_file), "Máy in", data_rows_1, start_row=2)
        
        # Write updated data with same key
        data_rows_2 = [
            ["001", "2026-01-15", "MachineA", "Line01", 5, "", "C001", "", "", "", "", "", "", "Updated", "", ""],
        ]
        stats = self.handler.write_to_summary_file(str(self.summary_file), "Máy in", data_rows_2, start_row=2)
        
        self.assertEqual(stats['inserted'], 0)
        self.assertEqual(stats['updated'], 1)
    
    def test_col_letter_to_idx(self):
        """Test column letter to index conversion."""
        self.assertEqual(ExcelHandler._col_letter_to_idx("A"), 1)
        self.assertEqual(ExcelHandler._col_letter_to_idx("B"), 2)
        self.assertEqual(ExcelHandler._col_letter_to_idx("Z"), 26)
        self.assertEqual(ExcelHandler._col_letter_to_idx("AA"), 27)
        self.assertEqual(ExcelHandler._col_letter_to_idx("P"), 16)
    
    def test_get_extended_data(self):
        """Test getting extended data as dictionaries."""
        headers, data_dicts = self.handler.get_extended_data(
            str(self.kdtps_file),
            "Máy in",
            start_row=3,
            header_row=2,
            end_col="P"
        )
        
        self.assertEqual(len(data_dicts), 3)
        self.assertIn('col_a', data_dicts[0])
        self.assertIn('col_n', data_dicts[0])
        self.assertEqual(data_dicts[0]['col_a'], "001")


class TestExcelHandlerSingleton(unittest.TestCase):
    """Test excel handler singleton."""
    
    def test_get_excel_handler_returns_same_instance(self):
        """Test singleton pattern."""
        handler1 = get_excel_handler()
        handler2 = get_excel_handler()
        self.assertIs(handler1, handler2)


if __name__ == '__main__':
    unittest.main()
